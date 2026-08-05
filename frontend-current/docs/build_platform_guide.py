from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("SFM_Compliance_Plattform_Vortragsleitfaden.docx")

BLUE = "2E74B5"
DARK = "0B1F3A"
TEAL = "007C89"
LIGHT_BLUE = "E8F4F8"
LIGHT_GRAY = "F4F7FA"
BORDER = "C9D6E2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    tbl = table._tbl
    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)
    tbl.insert(0, grid)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text, bold=False, color=None, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("SFM Compliance | Vortragsleitfaden")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7C8F")

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Plattform unterstützt Vorbereitung und Nachweisstrukturierung. Fachliche Freigabe bleibt beim Berater/Admin.")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7C8F")


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("SFM Compliance")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Vortragsleitfaden zur ISMS-/ISO-27001-/NIS-2-Plattform")
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("4A5D70")

    meta = doc.add_table(rows=1, cols=3)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(meta, [2.1, 2.1, 2.1])
    labels = [
        ("Zweck", "kurze Plattformvorstellung"),
        ("Stand", "16.06.2026"),
        ("Zielgruppe", "Kollegen, Kunden, Berater"),
    ]
    for cell, (label, value) in zip(meta.rows[0].cells, labels):
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text(cell, f"{label}\n{value}", bold=False, size=9.3)
    set_table_width(meta, [2.1, 2.1, 2.1])

    doc.add_paragraph()
    add_callout(
        doc,
        "Kernaussage",
        "Die Plattform ersetzt nicht den Berater. Sie führt Kunden durch Vorbereitung, Datenerfassung und Nachweise, macht formale Lücken sichtbar und bereitet die fachliche Prüfung vor. Bewertung, Freigabe und Auditbereitschaft bleiben immer beim Berater/Admin.",
    )


def add_callout(doc, title, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, "A7C9D8", "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(body)
    run.font.name = "Calibri"
    run.font.size = Pt(9.6)
    run.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE if level <= 2 else "1F4D78")
    run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 11.5)


def add_intro(doc):
    add_heading(doc, "Wie du die Plattform in 5 bis 7 Minuten vorstellst", 1)
    steps = [
        ("1. Start im Dashboard", "Zeige Gesamtstatus, offene Aufgaben, Nachweise, Review-Status und Blocker."),
        ("2. Kundenmodus zeigen", "Erkläre: Der Kunde sieht nicht die ganze ISO-Komplexität, sondern den nächsten einfachen Schritt."),
        ("3. Projektstrukturplan öffnen", "Zeige Arbeitspakete, Verantwortliche, Termine, Fortschritt und verlinkte Dokumente."),
        ("4. Compliance-Bereiche erklären", "ISO 27001 und NIS-2 werden getrennt geführt, aber gemeinsame Nachweise werden nutzbar gemacht."),
        ("5. Review Center zeigen", "Der Berater sieht Einreichungen, fordert Nacharbeit an und gibt fachlich frei."),
        ("6. Auditpaket-Gate erklären", "Die Plattform bereitet vor, behauptet aber nicht automatisch Auditbereitschaft."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [0.45, 2.1, 4.0]
    set_table_width(table, widths)
    headers = ["#", "Station", "Sprechpunkt"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "E8EEF5")
        set_cell_border(cell)
        set_cell_text(cell, text, bold=True, color=DARK, size=9.5)
    keep_row_together(table.rows[0])
    for idx, (station, speech) in enumerate(steps, start=1):
        row = table.add_row()
        keep_row_together(row)
        values = [str(idx), station, speech]
        for cell, text in zip(row.cells, values):
            set_cell_border(cell)
            set_cell_text(cell, text, bold=False, size=9.2)
    set_table_width(table, widths)
    doc.add_paragraph()


def add_group_table(doc, group_title, group_intro, rows):
    add_heading(doc, group_title, 1)
    p = doc.add_paragraph(group_intro)
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string("4A5D70")

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [0.45, 1.45, 3.05, 1.55]
    set_table_width(table, widths)
    headers = ["#", "Menüpunkt", "Kurz erklären", "Beim Vorführen zeigen"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "E8EEF5")
        set_cell_border(cell)
        set_cell_text(cell, text, bold=True, color=DARK, size=9.2)
    keep_row_together(table.rows[0])
    for idx, row_data in enumerate(rows, start=1):
        row = table.add_row()
        keep_row_together(row)
        values = [str(idx), *row_data]
        for col_idx, (cell, text) in enumerate(zip(row.cells, values)):
            set_cell_border(cell)
            if col_idx == 1:
                set_cell_text(cell, text, bold=True, color=DARK, size=9)
            else:
                set_cell_text(cell, text, size=8.8)
    set_table_width(table, widths)
    doc.add_paragraph()


def add_closing(doc):
    add_heading(doc, "Kernsätze für die Präsentation", 1)
    phrases = [
        "SFM Compliance ist kein Ersatz für Beratung, sondern ein Arbeitsraum, der Vorbereitung und Prüfung sauber trennt.",
        "Der Kunde sieht einfache nächste Schritte; der Berater sieht fachliche Prüf- und Freigabeinformationen.",
        "Nachweise, Risiken, Policies und Aufgaben werden nicht verstreut, sondern in einer gemeinsamen Struktur geführt.",
        "Das Auditpaket wird vorbereitet, aber erst durch Berater/Admin fachlich freigegeben.",
    ]
    for phrase in phrases:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.1)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run("• ")
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        run = p.add_run(phrase)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(DARK)

    add_callout(
        doc,
        "Abschlussformulierung",
        "Kurz gesagt: Die Plattform macht Kundenarbeit einfacher, Beraterarbeit nachvollziehbarer und den Weg zum Auditpaket kontrollierbarer, ohne fachliche Verantwortung zu automatisieren.",
        fill="F6FBFD",
    )


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)

    add_header_footer(doc)
    add_title(doc)
    add_intro(doc)
    doc.add_page_break()

    add_group_table(
        doc,
        "1. Überblick",
        "Diese Bereiche sind der beste Einstieg: Sie erklären sofort, wo das Projekt steht und was der Kunde als Nächstes tun soll.",
        [
            ("Dashboard", "Zentrale Startseite mit Programmstatus, Top-Risiken, offenen Aufgaben, Nachweisen, Review-Status und Auditpaket-Blockern.", "Gesamtstatus und wichtigste Kacheln"),
            ("Benachrichtigungen", "Inbox für neue Einreichungen, Fristen, Rückfragen, Freigaben und Nacharbeit.", "Offene Hinweise und Rückgaben"),
            ("Kundenmodus", "Geführte Schritt-für-Schritt-Strecke in einfacher Sprache: Was muss ich jetzt ausfüllen, einreichen oder nacharbeiten?", "Heute erledigen, Nacharbeit, Abgabe-Check"),
            ("Projektstrukturplan", "Chronologische Arbeitspakete für ISO 27001 und NIS-2 mit Owner, Status, Termin und verknüpften Templates/Nachweisen.", "Phasen, Verantwortliche, Fortschritt"),
        ],
    )

    add_group_table(
        doc,
        "2. Compliance",
        "Hier liegen die fachlichen Sichten. ISO 27001 und NIS-2 bleiben getrennt, gemeinsame Nachweise können aber mehrfach genutzt werden.",
        [
            ("ISO 27001", "Clause- und Annex-A-orientierte Sicht mit Arbeitspaketen, Nachweisen, Umsetzungshinweisen und Status.", "Clause/Annex-A-Auswahl und Arbeitsartefakte"),
            ("NIS-2", "Separater Bereich für Betroffenheit, Pflichten, Nachweise und offene Punkte. Keine automatische Gleichsetzung mit ISO 27001.", "Rechtsregister, Auditmatrix, Projektplan"),
            ("Dokumentation", "Zentrales Nachweis-Center für Uploads, Links, Metadaten, Versionen und Prüfstatus.", "Nachweis anlegen und Status zeigen"),
            ("Beraterprüfung", "Review Center für Einreichungen, Nacharbeit, Freigaben, Ablehnungen und auditrelevante Markierungen.", "Prüfen, Kommentar an Kunde, Freigabe"),
            ("Auditmatrix", "Übersicht, welche Anforderungen welchen Nachweisen, Aufgaben und Prüfständen zugeordnet sind.", "Nachweis- und Statusspalten"),
            ("SoA Matrix", "Statement-of-Applicability-Sicht: Anwendbarkeit, Umsetzung, Begründung, Nachweise und Prüfung.", "SoA-Zeile mit verknüpften Nachweisen"),
            ("Templates", "Ausfüllbare Vorlagen, Policy-Entwürfe und Prozesspakete mit Erklärungen und Beispielen.", "Vorlage öffnen und ausfüllen"),
        ],
    )
    doc.add_page_break()

    add_group_table(
        doc,
        "3. Register",
        "Die Register sammeln strukturierte Kundendaten. Ziel ist nicht Bürokratie, sondern eine einheitliche Datenbasis für Nachweise, Risiken und Aufgaben.",
        [
            ("Asset Management", "Erfasst wichtige Systeme, Informationen und Verantwortliche als Grundlage für Schutzbedarf, Risiken und Nachweise.", "Asset anlegen, Owner setzen"),
            ("Risiko Management", "Einfacher Risikoworkshop: Was ist wichtig, was kann schiefgehen, was machen wir dagegen, wer ist verantwortlich?", "Risiko erfassen und Matrix zeigen"),
            ("Rechtsregister", "Sammelt relevante rechtliche oder regulatorische Themen und deren Bearbeitungsstand.", "Eintrag, Zuständigkeit, Nachweis"),
            ("Vendor Management", "Lieferanten erfassen, Kritikalität bewerten, Vertrags-/Nachweisprüfung planen und Review anstoßen.", "Lieferant, Nachweis, Beraterprüfung"),
        ],
    )

    add_group_table(
        doc,
        "4. Betrieb",
        "Diese Bereiche machen die Plattform im Alltag nutzbar: Policies, Vorfälle, Verträge, Schnittstellen und Aufgaben werden nicht isoliert geführt.",
        [
            ("Policy Management", "Policy-Pakete, Register, Versionen und Dokumentanker. Policies können geöffnet, bearbeitet und zur Prüfung gegeben werden.", "Policy öffnen und Erklärung zeigen"),
            ("Incident Management", "Sicherheitsereignisse strukturiert erfassen, bewerten und bei Bedarf an Berater/Admin eskalieren.", "Incident erfassen und Status zeigen"),
            ("Contract Management", "Verträge mit Lieferanten, Nachweisen, Fristen und Verantwortlichen verknüpfen.", "Vertrag/Review-Fälligkeit"),
            ("Integrationen", "Roadmap für spätere API-Anbindungen und Nachweisquellen. Aktuell als strukturierte Integrationsplanung.", "Integrationsstatus 0/60 erklären"),
            ("Task Management", "Aufgaben aus Projektplan, Kundenmodus, Nacharbeit und Beraterprüfung zentral nachverfolgen.", "Aufgaben nach Status und Owner"),
        ],
    )
    doc.add_page_break()

    add_group_table(
        doc,
        "5. Administration",
        "Administration ist für interne Steuerung, Betrieb, Sicherheit und Nachvollziehbarkeit gedacht. Kunden sollten hier nur sehen, was ihre Rolle braucht.",
        [
            ("Team & Einladungen", "Mitarbeiter einladen, Rollen zuweisen und Zusammenarbeit im Kundenprojekt ermöglichen.", "Einladung und Rollenlogik"),
            ("Mein Konto", "Profil, Session, Passwort-/MFA-Hinweise und persönliche Einstellungen.", "Konto- und Sicherheitsstatus"),
            ("Produktivbetrieb", "Checkliste für produktionsrelevante Themen wie Auth, Rollen, Dateiablage, Logging und Betrieb.", "Produktionsstatus und offene Punkte"),
            ("Betriebsmonitor", "Technischer Überblick zu Systemzustand, Laufzeit, Alerts und Wartungsaktionen.", "Health-/Betriebsstatus"),
            ("Backup & Recovery", "Backup- und Wiederherstellungsplanung mit Status, Snapshots und Betriebschecks.", "Backup-Readiness"),
            ("Datenlebenszyklus", "Regeln für Aufbewahrung, Löschung, Verantwortlichkeiten und Datenbestände.", "Lifecycle-Karten"),
            ("Audit Trail", "Nachvollziehbarkeit: Wer hat was geändert, eingereicht, freigegeben oder zurückgegeben.", "Änderungshistorie und Versionen"),
            ("Beraterbereich", "Interner Bereich für Generatoren, Beraterlogik und nicht-kundensichtbare Hilfsfunktionen.", "Nur Berater/Admin sichtbar"),
        ],
    )

    add_closing(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
