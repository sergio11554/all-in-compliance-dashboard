from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("SFM_Compliance_Ausfuehrliche_Funktionsuebersicht_2026-07-08.docx")

DARK = "0B1F3A"
BLUE = "2E74B5"
TEAL = "007C89"
MUTED = "4A5D70"
LIGHT_BLUE = "E8F4F8"
LIGHT_GRAY = "F2F4F7"
LIGHT_TEAL = "EEF9FA"
LIGHT_AMBER = "FFF6E6"
BORDER = "C9D6E2"


def twips(inches):
    return str(int(inches * 1440))


def de(text):
    repl = {
        "fuer": "für", "Fuer": "Für", "ueber": "über", "Ueber": "Über",
        "gefuehrt": "geführt", "Gefuehrt": "Geführt", "fuehrt": "führt",
        "Luecken": "Lücken", "luecken": "lücken", "Pruefung": "Prüfung",
        "pruefung": "prüfung", "Pruef": "Prüf", "pruef": "prüf",
        "Nachvollziehbarkeit": "Nachvollziehbarkeit", "koennen": "können",
        "Koennen": "Können", "muessen": "müssen", "Muessen": "Müssen",
        "ausfuell": "ausfüll", "Ausfuell": "Ausfüll", "oeffnen": "öffnen",
        "Oeffnen": "Öffnen", "naechst": "nächst", "Naechst": "Nächst",
        "noetig": "nötig", "Noetig": "Nötig", "Haertung": "Härtung",
        "laeuft": "läuft", "klaeren": "klären", "Klaeren": "Klären",
        "verknuepf": "verknüpf", "Verknuepf": "Verknüpf",
        "Loesch": "Lösch", "Quarantaene": "Quarantäne",
        "eigenstaendig": "eigenständig", "Vorgaenge": "Vorgänge",
        "Vorgaenge": "Vorgänge", "Rueck": "Rück", "gross": "groß",
        "Gross": "Groß", "erfuell": "erfüll", "zurueck": "zurück",
        "Qualitaet": "Qualität", "Vertraege": "Verträge",
        "Erklaer": "Erklär", "erklaer": "erklär",
    }
    for old, new in repl.items():
        text = text.replace(old, new)
    return text


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, val in [("top", 90), ("start", 120), ("bottom", 90), ("end", 120)]:
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def fixed_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), twips(width))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        keep_row(row)
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), twips(width))


def cell_text(cell, text, bold=False, color=DARK, size=9.2):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(de(text))
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def para(doc, text, size=10.4, color=MUTED, bold=False, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(de(text))
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(de(text), level=level)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(de(text))
    r.font.name = "Calibri"
    r.font.size = Pt(10.1)
    r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def callout(doc, title, body, fill=LIGHT_BLUE, border_color="A7C9D8"):
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [6.5])
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    border(cell, border_color, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(de(title))
    r.font.name = "Calibri"
    r.font.size = Pt(10.4)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(de(body))
    r.font.name = "Calibri"
    r.font.size = Pt(9.8)
    r.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph()


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    fixed_table(t, widths)
    for cell, text in zip(t.rows[0].cells, headers):
        shade(cell, LIGHT_GRAY)
        border(cell)
        cell_text(cell, text, bold=True, size=9.2)
    for values in rows:
        row = t.add_row()
        keep_row(row)
        for cell, text in zip(row.cells, values):
            border(cell)
            cell_text(cell, text, size=8.9)
    doc.add_paragraph()
    return t


def header_footer(doc):
    header = doc.sections[0].header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("SFM Compliance | ausführliche Funktionsübersicht")
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("6B7C8F")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Arbeitsstand der Plattform; finale fachliche Bewertung bleibt beim Berater/Admin.")
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("6B7C8F")


def add_cover(doc):
    para(doc, "SFM Compliance", size=26, color=DARK, bold=True, after=2)
    para(doc, "Ausführliche Übersicht: Was auf der Plattform bereits umgesetzt ist", size=15, color=MUTED, after=14)
    meta = doc.add_table(rows=1, cols=3)
    fixed_table(meta, [2.1, 2.1, 2.3])
    for cell, text in zip(meta.rows[0].cells, ["Stand\n08.07.2026", "Zweck\nKollegen vollständig abholen", "Fokus\nUmgesetzte Funktionen + Technik"]):
        shade(cell, LIGHT_GRAY)
        border(cell)
        cell_text(cell, text, size=9.5)
    doc.add_paragraph()
    callout(
        doc,
        "Worum es in diesem Dokument geht",
        "Dieses Dokument beschreibt konkret, welche Funktionen in SFM Compliance bereits gebaut wurden: Oberfläche, Kundenmodus, Compliance-Bereiche, Register, Review Center, Auditpaket, Versionierung, Backend, Rollen, Mandanten, Uploads, Betrieb und Deployment-Bausteine.",
    )
    callout(
        doc,
        "Produktprinzip",
        "Die Plattform ersetzt nicht den Berater. Sie führt Kunden durch die Vorbereitung, sammelt Informationen und Nachweise, macht formale Lücken sichtbar und bereitet die fachliche Prüfung vor. Freigabe, Ablehnung, Audit-Relevanz und Auditbereitschaft bleiben beim Berater/Admin.",
        fill=LIGHT_AMBER,
        border_color="E8C06F",
    )


def add_summary(doc):
    heading(doc, "1. Kurzüberblick: Was wurde insgesamt aufgebaut?", 1)
    for item in [
        "Eine zentrale Webplattform für ISMS-, ISO-27001-, NIS-2- und Compliance-Vorbereitung unter dem Produktnamen SFM Compliance.",
        "Ein geführter Kundenmodus, der Kunden in einfacher Sprache durch Startprofil, Assets, Risiken, Policies, Nachweise, SoA und Abgabecheck führt.",
        "Ein Beraterbereich mit Review Center, in dem Einreichungen geprüft, freigegeben, abgelehnt oder zur Nacharbeit zurückgegeben werden können.",
        "Ein Dashboard/Tool-Hub für Management, Kunden und Berater mit nächsten Schritten, Status, Nachweisen, Risiken, NIS-2-Navigation und Auditpaket.",
        "Mehrere operative Register: Assets, Risiken, Rechtsregister, Lieferanten, Policies, Incidents, Verträge und Aufgaben.",
        "Ein Dokumentations- und Nachweisbereich mit Upload, Metadaten, Zuordnung, Einreichstatus, Versionierung und Audit-Trail.",
        "Eine SoA-Matrix und Auditmatrix, um ISO-27001-Controls, Nachweise, Lücken und Prüfstatus strukturiert sichtbar zu machen.",
        "Ein Backend mit Authentifizierung, Sessions, Rollen, Mandanten, Dateiablage, Audit-Log, Backups, Benachrichtigungen, Hintergrundjobs und Healthcheck.",
        "Deployment- und Betriebsunterlagen wie Docker, systemd, nginx, Go-Live-Checkliste, Backup-/Restore-Runbook und Operations-Handbuch.",
        "Ein sukzessiv überarbeitetes Enterprise-SaaS-Design mit ruhigerer Sidebar, konsistenten Buttons, Badges, Tabellen, Panels und Widescreen-Layouts.",
    ]:
        bullet(doc, item)


def add_frontend(doc):
    heading(doc, "2. Umgesetzte Oberfläche und Nutzerführung", 1)
    rows = [
        ("Dashboard / Tool-Hub", "Zentrale Startseite mit wichtigsten Einstiegspunkten: Heute erledigen, Beraterprüfung, Nachweise, Risiken, NIS-2 Navigator und Auditpaket.", "Kunde und Berater sehen sofort, wo sie weitermachen müssen."),
        ("Management-Ansicht", "KPI- und Statusbereiche für Gesamtstand, Top-Risiken, offene Entscheidungen, Fristen und Compliance-Spuren.", "Kollegen können zeigen, dass die Plattform nicht nur Datenerfassung ist, sondern auch Steuerung."),
        ("Kundenmodus", "Schritt-für-Schritt-Strecke mit einfacher Sprache, Nacharbeit, Aufgaben, Einreichstatus und Übergabecheck.", "Nicht-IT-Kunden sollen verstehen, was sie konkret tun müssen."),
        ("Benachrichtigungen", "Zentrale Inbox für Hinweise, Rückfragen, Eskalationen und Systemmeldungen mit Filter und Aktionen.", "Wichtige Rückfragen landen nicht verteilt in E-Mails."),
        ("Projektstrukturplan", "Chronologische Arbeitspakete für ISO 27001 und NIS-2, inklusive Owner, Fortschritt, Status, Nachweise und Verlinkungen.", "Das Projekt wird als Ablauf sichtbar, nicht als unübersichtliche Dokumentensammlung."),
        ("Hilfetexte und Beispiele", "Felder und Vorlagen enthalten erklärende Hinweise und Beispiele, was einzutragen ist.", "Kunden können besser selbst ausfüllen; Rückfragen sinken."),
    ]
    table(doc, ["Bereich", "Was wurde umgesetzt?", "Nutzen"], rows, [1.55, 3.25, 1.70])


def add_compliance(doc):
    heading(doc, "3. Umgesetzte Compliance-Bereiche", 1)
    rows = [
        ("ISO 27001", "Eigener ISO-Bereich mit Clauses, Annex-A-nahen Controls, Arbeitsartefakten, Nachweisverknüpfung und Registerabsprung.", "ISO bleibt klar getrennt und strukturiert."),
        ("NIS-2", "Eigener NIS-2 Navigator mit Einordnung, Pflichten, offenen Punkten, Nachweisen und Auditmatrix-/Projektplanbezug.", "NIS-2 wird nicht einfach als ISO-Anhang behandelt."),
        ("SoA Matrix", "Statement of Applicability als tabellarische Arbeitsansicht mit Controls, Status, Nachweisen, Bewertung und Beraterprüfung.", "SoA kann direkt in der Plattform gepflegt werden."),
        ("Auditmatrix", "Prüf-/Nachweismatrix mit Status, offenen Lücken und Verbindung zu Arbeitsbereichen.", "Berater sehen, wo Nachweise fehlen oder ungeprüft sind."),
        ("Arbeitsvorlagen", "Template-Bibliothek mit ausfüllbaren Vorlagen, Hilfetexten, Beispielen, Quellen-/Kontextinfos und direkter Öffnen-Funktion.", "Kunden bekommen nicht nur Stichpunkte, sondern ausfüllbare Dokumentationsanker."),
        ("Policy Management", "Policy-Paket mit ISO-Kern-Policies, NIS-2-Überschneidungen, Annex-A-/ISMS-Verfahren, Register und Vorlagenöffnung.", "Policies werden als steuerbarer Prozess geführt."),
    ]
    table(doc, ["Modul", "Umsetzung", "Bedeutung"], rows, [1.45, 3.35, 1.70])


def add_registers(doc):
    heading(doc, "4. Umgesetzte Register und operative Module", 1)
    rows = [
        ("Asset Management", "Erfassung von Systemen, Informationen, Prozessen, Ownern, Zuordnung, Schutz-/Nachweisbedarf und Beraterstatus.", "Grundlage für Risiken, Nachweise und Projektarbeit."),
        ("Risikomanagement", "Risikoworkshop, Risikoregister, Methodik, Bewertung, Behandlung, Reststatus, Akzeptanz, Wirksamkeit und Traceability.", "Risiken werden verständlich und nachvollziehbar bearbeitet."),
        ("Rechtsregister", "Pflichten/Regelungen mit Verantwortung, Status, Frist, Nachweis und Review.", "Regulatorische Arbeit wird als eigenes Register steuerbar."),
        ("Vendor Management", "Lieferantenregister mit Kritikalität, Review, Nachweisen, Vertragsbezug und Beraterprüfung.", "Lieferkette/Supplier Security wird praktisch abbildbar."),
        ("Incident Management", "Erfassung von Ereignissen/Vorfällen, Bewertung, Aufgaben, Status, Nachweise und Lessons-Learned-Logik.", "Vorfälle können strukturiert dokumentiert werden."),
        ("Contract Management", "Verträge, Owner, Review-Termine, Status, Nachweise und Verknüpfung mit Aufgaben.", "Verträge werden in Compliance-Prozesse eingebunden."),
        ("Task Management", "Aufgaben, Nacharbeit, offene Punkte, Fristen, Owner, Prüfschritte und Nachweisbezug.", "Aus Plattformdaten entstehen konkrete To-dos."),
        ("Integrationen", "Integrationskatalog/Ansicht für geplante und aktive Schnittstellen mit Status und Nachweisen.", "API-/Tool-Anbindungen können vorbereitet und transparent geführt werden."),
    ]
    table(doc, ["Modul", "Was kann es bereits?", "Warum wichtig?"], rows, [1.45, 3.35, 1.70])


def add_review_audit(doc):
    heading(doc, "5. Beraterprüfung, Auditpaket und Nachvollziehbarkeit", 1)
    for item in [
        "Review Center: zentrale Liste für eingereichte Nachweise, Risiken, SoA-Einträge, Policies, Assets, Aufgaben und auditpaket-relevante Punkte.",
        "Statuslogik: offen, vom Kunden eingereicht, in Beraterprüfung, Nacharbeit nötig, freigegeben, abgelehnt und auditrelevant.",
        "Nur Berater/Admin dürfen final freigeben, ablehnen oder auditrelevant markieren; Kunden dürfen einreichen und nacharbeiten.",
        "Detailprüfung mit Kundeneingabe, verknüpftem Nachweis, ISO-/NIS-2-Zuordnung, internem Beraterkommentar und Kommentar an Kunde.",
        "Nacharbeit erzeugt Kundenaufgaben, die im Kundenmodus sichtbar werden.",
        "Auditpaket-Gate unterscheidet formal vorbereitet von beraterfreigegeben.",
        "Audit Trail und Versionierung dokumentieren Änderungen, Einreichungen, Freigaben, Rückgaben und alte Versionen.",
        "Dokumenten- und Policy-Versionen können nachvollziehbar gehalten werden.",
    ]:
        bullet(doc, item)


def add_backend(doc):
    heading(doc, "6. Technischer Hintergrund: Was läuft bereits im Backend?", 1)
    rows = [
        ("Server / Healthcheck", "Lokaler Backend-Server läuft über http://127.0.0.1:5173; /api/health liefert OK.", "Die Plattform ist nicht nur file:// oder statisches HTML."),
        ("Auth / Sessions", "Login, Logout, Sessionprüfung, eigene Sessions listen und Sessions widerrufen.", "Grundlage für geschützte Kundenräume."),
        ("Rollenrechte", "Admin, Consultant, Manager, Customer, Contributor, Auditor und Viewer mit abgestuften Rechten.", "Kunden können arbeiten, aber nicht final freigeben."),
        ("MFA", "MFA-Setup, Aktivierung und Deaktivierung sind als Backend-Endpunkte vorhanden.", "Wichtiger Security-Baustein für Admin/Manager."),
        ("Mandanten", "Mandantenübersicht, Mandantenanlage, Update, Export, Lifecycle und Purge.", "Grundlage für mehrere Kundenprojekte/Kundenräume."),
        ("Workspace State", "State laden/speichern, Metadaten, Snapshots erstellen und wiederherstellen.", "Arbeitsstand kann serverseitig verwaltet werden."),
        ("Dateien", "Upload, Download, Dateiliste, Quarantäne, Rescan und kontrolliertes Löschen isolierter Dateien.", "Nachweise werden technisch verwaltbar."),
        ("Audit Log", "Audit-Log anzeigen, Integrität prüfen und Export erzeugen.", "Wer hat was wann gemacht, ist nachvollziehbar."),
        ("Backups", "Backup erstellen, Backup-Liste, Backup-Download und Restore-/Runbook-Unterlagen.", "Betriebs- und Wiederherstellungsfähigkeit vorbereitet."),
        ("Jobs / Operations", "Hintergrundjobs, Retry/Cancel, Operations Alerts, Wartungsstatus und Betriebsmonitoring.", "Längere Aufgaben und technische Hinweise werden steuerbar."),
        ("Benachrichtigungen", "Benachrichtigungsstatus, Zustellqueue, Retry, Webhook/SMTP-Konfiguration in Operations-Unterlagen.", "Systemmeldungen können zentral verarbeitet werden."),
    ]
    table(doc, ["Backend-Baustein", "Umgesetzter Stand", "Einordnung"], rows, [1.45, 3.25, 1.80])


def add_admin_ops(doc):
    heading(doc, "7. Administration, Team und Produktivbetrieb", 1)
    rows = [
        ("Team & Einladungen", "Mitarbeiter können vorgesehen in Kundenräume eingeladen und Rollen zugewiesen werden.", "Kunden können mit mehreren Beteiligten arbeiten."),
        ("Mein Konto", "Konto-/Sessionbereich mit Passwort-/MFA- und Sicherheitsstatus.", "Benutzer können eigene Sicherheit verwalten."),
        ("Produktivbetrieb", "Checklisten- und Betriebsbereiche für technische Reife, Security, Backups, Jobs, Quarantäne, Lifecycle und Operations.", "Plattform denkt Betrieb mit, nicht nur UI."),
        ("Deployment", "Dockerfile, docker-compose.yml, systemd-Services, nginx-Konfiguration und PostgreSQL-Schema liegen vor.", "Basis für echten Serverbetrieb ist vorbereitet."),
        ("Betriebsdokumente", "OPERATIONS.md, Deployment-Quickstart, Go-Live-Checkliste, Backup-/Restore-Runbook, Security-Hardening-Checkliste, UX-Abnahmeplan.", "Kollegen sehen, dass Betrieb und Übergabe bedacht wurden."),
    ]
    table(doc, ["Bereich", "Bereits vorhanden", "Nutzen"], rows, [1.45, 3.35, 1.70])


def add_design(doc):
    heading(doc, "8. Design- und UX-Arbeit, die bereits umgesetzt wurde", 1)
    for item in [
        "Produktname SFM Compliance wurde in der Oberfläche verankert.",
        "Das Design wurde mehrfach Richtung Enterprise SaaS überarbeitet: ruhigere Farben, weniger Badge-Lärm, konsistentere Buttons und Tabellen.",
        "Sidebar wurde strukturiert: Start, Prüfung, Compliance, Register, Betrieb, Administration und weitere Werkzeuge.",
        "Widescreen-Layouts wurden stabilisiert, damit große Bildschirme nicht leer oder gequetscht wirken.",
        "Formulare, Registerkarten, Review-Listen, Tabellen und Status-Pills wurden vereinheitlicht.",
        "Kundenansichten wurden vereinfacht: weniger Fachsprache, mehr klare nächste Aktionen.",
        "Hilfetexte und Beispiele wurden bewusst beibehalten, damit Nutzer wissen, was sie eintragen sollen.",
        "Die UI vermeidet automatische Aussagen wie ISO-konform, NIS-2-erfüllt oder auditbereit.",
    ]:
        bullet(doc, item)


def add_boundaries(doc):
    heading(doc, "9. Was die Plattform bewusst noch nicht behauptet", 1)
    callout(
        doc,
        "Wichtige Grenze",
        "Die Plattform zeigt Vorbereitungsstand, formale Lücken, Einreichungen und Review-Status. Sie ersetzt keine fachliche Bewertung, keine rechtliche Einordnung und keine Auditentscheidung.",
        fill=LIGHT_AMBER,
        border_color="E8C06F",
    )
    for item in [
        "Keine automatische Aussage: Der Kunde ist ISO-konform.",
        "Keine automatische Aussage: NIS-2 ist vollständig erfüllt.",
        "Keine automatische Aussage: Das Audit ist bestanden oder auditbereit.",
        "Automatische Checks dürfen höchstens formal vollständig, Nachweis vorhanden oder zur Prüfung bereit anzeigen.",
        "Fachliche Freigaben und Auditpaket-Freigaben bleiben beim Berater/Admin.",
    ]:
        bullet(doc, item)


def add_next(doc):
    heading(doc, "10. Nächste große Ausbauschritte", 1)
    rows = [
        ("Echte Pilotdaten testen", "Mit einem realistischen Beispielkunden alle Module einmal vollständig durchspielen.", "hoch"),
        ("QA und Design-Abnahme", "Alle Hauptseiten Desktop/Widescreen/Mobile prüfen; keine Überläufe, keine abgeschnittenen Texte.", "hoch"),
        ("Produktiv-Security", "Secrets, Cookies, HSTS, MFA-Pflicht, Upload-Scanner, Rechte und Session-Limits produktionsnah härten.", "hoch"),
        ("Deployment finalisieren", "PostgreSQL/S3/Reverse Proxy/Docker/systemd testen und dokumentiert betreiben.", "hoch"),
        ("Datenschutz/Recht", "AVV, Rollenverantwortung, Löschkonzept, Aufbewahrung und Kundendatenprozess intern freigeben.", "hoch"),
        ("Content-Freigabe", "Policy-/Template-Texte, Beispiele und fachliche Hilfetexte final prüfen.", "mittel"),
        ("Demo-Paket", "Screenshots, kurzer Pitch, ausführliche Funktionsübersicht und Live-Demo-Reihenfolge final zusammenstellen.", "mittel"),
    ]
    table(doc, ["Schritt", "Was noch zu tun ist", "Priorität"], rows, [1.55, 4.05, 0.90])


def build():
    doc = Document()
    style_doc(doc)
    header_footer(doc)
    add_cover(doc)
    add_summary(doc)
    add_frontend(doc)
    add_compliance(doc)
    add_registers(doc)
    add_review_audit(doc)
    add_backend(doc)
    add_admin_ops(doc)
    add_design(doc)
    add_boundaries(doc)
    add_next(doc)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
