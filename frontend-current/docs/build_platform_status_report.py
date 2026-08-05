from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("SFM_Compliance_Standbericht_Kollegen_2026-07-08.docx")

DARK = "0B1F3A"
BLUE = "2E74B5"
TEAL = "007C89"
MUTED = "4A5D70"
LIGHT_BLUE = "E8F4F8"
LIGHT_GRAY = "F2F4F7"
LIGHT_TEAL = "EEF9FA"
LIGHT_AMBER = "FFF6E6"
LIGHT_GREEN = "F0FAF4"
BORDER = "C9D6E2"
WHITE = "FFFFFF"


def de(text):
    replacements = {
        "fuer": "für",
        "Fuer": "Für",
        "gefuehr": "geführ",
        "Gefuehr": "Geführ",
        "Fuehr": "Führ",
        "fuehr": "führ",
        "Lueck": "Lück",
        "lueck": "lück",
        "Pruef": "Prüf",
        "pruef": "prüf",
        "pruefung": "prüfung",
        "erfuell": "erfüll",
        "ausfuell": "ausfüll",
        "Ausfuell": "Ausfüll",
        "oeffnen": "öffnen",
        "Oeffnen": "Öffnen",
        "koennen": "können",
        "Koennen": "Können",
        "muessen": "müssen",
        "Muessen": "Müssen",
        "naechst": "nächst",
        "Naechst": "Nächst",
        "noetig": "nötig",
        "Noetig": "Nötig",
        "gross": "groß",
        "Gross": "Groß",
        "Haertung": "Härtung",
        "haeng": "häng",
        "Qualitaet": "Qualität",
        "Quarantaene": "Quarantäne",
        "Vertraege": "Verträge",
        "eigenstaendig": "eigenständig",
        "Vorfalle": "Vorfälle",
        "Mandantenraeume": "Mandantenräume",
        "Rueckgaben": "Rückgaben",
        "ueber": "über",
        "Ueber": "Über",
        "zurueck": "zurück",
        "Zurueck": "Zurück",
        "Loesch": "Lösch",
        "Loeschen": "Löschen",
        "pruefen": "prüfen",
        "angrenzende": "angrenzende",
        "klaer": "klär",
        "Klaer": "Klär",
        "laeuft": "läuft",
        "Laeuft": "Läuft",
        "verknuepf": "verknüpf",
        "Verknuepf": "Verknüpf",
        "erklaer": "erklär",
        "Erklaer": "Erklär",
        "steuern": "steuern",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def twips(inches):
    return str(int(inches * 1440))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def fixed_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), twips(width))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        keep_row(row)
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), twips(width))


def cell_text(cell, text, bold=False, color=DARK, size=9.4, align=None):
    text = de(text)
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(de("SFM Compliance | Plattform-Standbericht"))
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7C8F")

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(de("Arbeitsstand fuer interne Vorstellung; fachliche Freigabe und Auditbereitschaft bleiben beim Berater/Admin."))
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7C8F")


def para(doc, text="", size=10.4, color=MUTED, bold=False, after=6, before=0, align=None):
    text = de(text)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(de(text), level=level)


def bullet(doc, text, level=0, bold_prefix=None):
    text = de(text)
    if bold_prefix:
        bold_prefix = de(bold_prefix)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.30 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.font.name = "Calibri"
        r.font.size = Pt(10.2)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(DARK)
        text = text[len(bold_prefix):]
        r = p.add_run(text)
    else:
        r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.2)
    r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def number(doc, text):
    text = de(text)
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.2)
    r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def callout(doc, title, body, fill=LIGHT_BLUE, border_color="A7C9D8"):
    title = de(title)
    body = de(body)
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [6.5])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_border(cell, border_color, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.name = "Calibri"
    r.font.size = Pt(10.3)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    r.font.name = "Calibri"
    r.font.size = Pt(9.8)
    r.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph()


def make_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    fixed_table(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        cell_text(cell, text, bold=True, size=9.2)
    for values in rows:
        row = table.add_row()
        keep_row(row)
        for cell, text in zip(row.cells, values):
            set_cell_border(cell)
            cell_text(cell, text, size=9.0)
    doc.add_paragraph()
    return table


def add_cover(doc):
    para(doc, "SFM Compliance", size=26, color=DARK, bold=True, after=2)
    para(doc, "Standbericht fuer Kollegen und interne Vorstellung", size=15, color=MUTED, after=14)

    meta = doc.add_table(rows=1, cols=3)
    fixed_table(meta, [2.1, 2.1, 2.3])
    for cell, text in zip(
        meta.rows[0].cells,
        [
            "Stand\n08.07.2026",
            "Format\nVortragsbericht + Handout",
            "Ziel\nKollegen in 15-30 Minuten abholen",
        ],
    ):
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_border(cell)
        cell_text(cell, text, size=9.5)

    doc.add_paragraph()
    callout(
        doc,
        "Kernaussage fuer die Vorstellung",
        "Die Plattform ist ein gefuehrter ISMS-/Compliance-Workspace. Sie ersetzt nicht den Berater, sondern sammelt Kundendaten, strukturiert Nachweise, macht formale Luecken sichtbar und bereitet die fachliche Beraterpruefung vor.",
    )
    callout(
        doc,
        "Abgrenzung",
        "Die Plattform darf nicht automatisch behaupten, ein Kunde sei ISO-konform, NIS-2-erfuellt oder auditbereit. Sie kann den Vorbereitungsstand zeigen; fachliche Freigabe, Ablehnung, Audit-Relevanz und Auditpaket-Freigabe bleiben beim Berater/Admin.",
        fill=LIGHT_AMBER,
        border_color="E8C06F",
    )


def add_talk_track(doc):
    heading(doc, "1. Sprechleitfaden fuer die ersten Minuten", 1)
    para(
        doc,
        "Diese Kurzfassung eignet sich als Einstieg, bevor die Plattform live gezeigt wird.",
        color=MUTED,
    )
    for item in [
        "Wir bauen SFM Compliance als Arbeitsplattform fuer ISO 27001, NIS-2 und angrenzende Compliance-Vorbereitung.",
        "Der Kunde sieht nicht zuerst Normsprache, sondern klare naechste Schritte: Was muss ich ausfuellen, hochladen oder klaeren?",
        "Der Berater sieht zentral, was eingereicht wurde, wo Nacharbeit noetig ist und welche Punkte ein Auditpaket blockieren.",
        "ISO 27001 und NIS-2 werden in der UI getrennt gefuehrt, gemeinsame Daten wie Assets, Risiken, Lieferanten und Nachweise werden aber mehrfach nutzbar gemacht.",
        "Im Hintergrund laeuft ein Backend mit Rollen, Sessions, Uploads, Mandantenlogik, Audit-Log, Backups und Betriebsmonitoring.",
        "Der aktuelle Stand ist vorfuehrbar und funktional breit, aber fuer echten Kundenbetrieb brauchen wir weiterhin Security-/Betriebs-Haertung, QA, Datenschutzpruefung und Deployment-Freigabe.",
    ]:
        bullet(doc, item)


def add_function_map(doc):
    heading(doc, "2. Was die Plattform aktuell fachlich abdeckt", 1)
    rows = [
        ("Dashboard / Tool-Hub", "Kuratierte Startseite mit naechsten Aufgaben, Nachweisen, Review, Risiken, NIS-2 und Auditpaket.", "Sofort verstehen: Was ist wichtig, was fehlt, was wartet?"),
        ("Kundenmodus", "Gefuehrte Strecke von Startprofil ueber Assets, Risiken, Policies, Nachweise, SoA bis Abgabecheck.", "Kunde muss weniger Normsprache verstehen und sieht konkrete Schritte."),
        ("Projektstrukturplan", "Chronologische Arbeitspakete, Owner, Status, Termine, Fortschritt und Dokument-/Template-Verlinkung.", "Projekt wird steuerbar statt Dokumentenablage."),
        ("ISO 27001", "Strukturierte Sicht auf Clauses, Annex-A-nahe Controls, Nachweise, Arbeitsartefakte und SoA-Bezug.", "ISO bleibt separat sichtbar und kann mit Nachweisen verbunden werden."),
        ("NIS-2 Navigator", "Separater NIS-2-Bereich fuer Einordnung, Pflichten, offene Punkte und Nachweisbezug.", "NIS-2 wird nicht unkontrolliert mit ISO vermischt."),
        ("Dokumentation / Nachweise", "Upload, Metadaten, Zuordnung, Versionen, Pruefstatus und Nachweis-Gaps.", "Nachweise sind auffindbar und pruefbar."),
        ("Beraterpruefung / Review Center", "Zentrale Sicht auf Einreichungen, Nacharbeit, Freigaben, Ablehnungen und Audit-Blocker.", "Fachliche Entscheidung bleibt beim Berater/Admin."),
        ("Auditpaket", "Auditpaket kann formal vorbereitet und gegen Gate-Kriterien geprueft werden.", "Exportbereit erst nach Beraterfreigabe, nicht automatisch auditbereit."),
        ("SoA Matrix", "Statement of Applicability als strukturierte Tabelle mit Status, Nachweisen und Pruefhinweisen.", "SoA wird als Arbeitsinstrument statt statischem Excel gedacht."),
        ("Auditmatrix", "Audit-/Nachweismatrix mit Luecken, Status und Verbindung zu Arbeitsbereichen.", "Berater sieht, welche Nachweise noch fehlen oder ungeprueft sind."),
    ]
    make_table(doc, ["Bereich", "Aktueller Funktionsstand", "Nutzen im Vortrag"], rows, [1.55, 2.75, 2.20])


def add_registers(doc):
    heading(doc, "3. Register und operative Arbeitsbereiche", 1)
    rows = [
        ("Asset Management", "Assets, Systeme, Informationen, Owner, Zuordnung ISO/NIS-2, Nachweisbedarf.", "Gemeinsame Grundlage fuer Risiken, Nachweise und Projektarbeit."),
        ("Risikomanagement", "Risiken erfassen, bewerten, Behandlung dokumentieren, Workshop-Fragen und Registeransicht.", "Risiken werden kundennah abgefragt: Was ist wichtig, was kann schiefgehen, was tun wir dagegen?"),
        ("Rechtsregister", "Rechts-/Pflichtenpunkte mit Verantwortung, Status, Nachweis und Review.", "Regulatorische Punkte werden nicht in ISO versteckt."),
        ("Vendor Management", "Lieferanten, Kritikalitaet, Review, Nachweise, Vertragsbezug und Beraterpruefung.", "Lieferkette wird als eigenstaendiger Arbeitsbereich steuerbar."),
        ("Policy Management", "Policy-Paket, ausfuellbare Vorlagen, Vorlagenoeffnung, Beispiele, Status und Freigabeprozess.", "Kunde kann aus Vorlagen arbeiten; Berater prueft final."),
        ("Incident Management", "Sicherheitsereignisse erfassen, bewerten, Aufgaben/Nachweise verbinden.", "Vorfalle werden strukturiert dokumentiert und eskalierbar."),
        ("Contract Management", "Vertraege, Owner, Review-Termine, Nachweise und Status.", "Vertragsbezogene Pflichten bleiben nicht verteilt in Dateien."),
        ("Task Management", "Offene Aufgaben, Nacharbeit, Fristen, Verantwortung und Pruefstatus.", "Kunde sieht konkret, was als Naechstes zu tun ist."),
    ]
    make_table(doc, ["Arbeitsbereich", "Was ist vorhanden?", "Warum relevant?"], rows, [1.55, 2.70, 2.25])


def add_consultant_flow(doc):
    heading(doc, "4. Beratersteuerung: Was wurde bewusst eingebaut?", 1)
    callout(
        doc,
        "Wichtig fuer die Positionierung",
        "Die Plattform ist kein automatischer Auditor. Sie trennt Vorbereitungslogik von fachlicher Freigabe. Kunden duerfen einreichen und nacharbeiten; Berater/Admin duerfen pruefen, freigeben, ablehnen und auditrelevante Punkte markieren.",
        fill=LIGHT_TEAL,
    )
    for item in [
        "Review Center mit neuen Einreichungen, Nacharbeit, Freigaben und Audit-Blockern.",
        "Detailpanel fuer Kundeneingabe, verknuepften Nachweis, ISO-/NIS-2-Zuordnung und Kommentare.",
        "Zwei Kommentararten: interner Beraterkommentar und Kommentar an den Kunden.",
        "Nacharbeit erzeugt Kundenaufgaben, damit Rueckgaben nicht in E-Mails verloren gehen.",
        "Auditpaket-Gate unterscheidet formal vorbereitet von beraterfreigegeben.",
        "Statuslogik: offen, vom Kunden eingereicht, in Beraterpruefung, Nacharbeit noetig, freigegeben, abgelehnt, auditrelevant.",
    ]:
        bullet(doc, item)


def add_technical(doc):
    heading(doc, "5. Technischer Hintergrund und Backend-Stand", 1)
    rows = [
        ("Backend-Modus", "Lokaler Backend-Server laeuft ueber http://127.0.0.1:5173; Healthcheck ist vorhanden und aktuell erfolgreich.", "Vortrag: Es ist nicht nur eine statische HTML-Demo."),
        ("Session / Auth", "Login, Session-Endpunkte, Rollen, MFA-Setup/Enable/Disable und Session-Revoke sind im Backend vorgesehen.", "Grundlage fuer Kunden- und Beraterzugriffe."),
        ("Rollenrechte", "Admin, Consultant, Manager, Customer, Contributor, Auditor und Viewer mit abgestuften Berechtigungen.", "Kunden koennen einreichen, aber nicht final freigeben."),
        ("Mandanten", "Mandanten-/Tenant-Endpunkte, Tenant-Overview, Lifecycle, Export und Purge sind vorhanden.", "Basis fuer mehrere Kundenraeume."),
        ("Workspace State", "State lesen/speichern, Snapshots erstellen und wiederherstellen, State-Metadaten.", "Nachvollziehbarer Arbeitsstand statt reiner Browser-Speicher."),
        ("Dateien / Uploads", "Dateiliste, Upload, Download, Quarantaene, Rescan und Metadaten.", "Nachweise koennen technisch abgelegt und kontrolliert werden."),
        ("Audit Log", "Audit-Log, Integritaetscheck und Export-Endpunkt.", "Wer hat was wann gemacht, wird technisch nachvollziehbar."),
        ("Backups", "Backup-Liste, Backup-Erstellung und Download plus Runbook/Deployment-Unterlagen.", "Wichtiger Baustein fuer Betrieb und Wiederherstellung."),
        ("Operations / Jobs", "Betriebsalarme, Hintergrundjobs, Retry/Cancel, Job-Artefakte.", "Langlaufende Aufgaben wie Backups oder Pruefungen muessen nicht im UI haengen."),
        ("Notifications", "Benachrichtigungsstatus und optionale Zustellqueue/Webhook/SMTP-Konfiguration.", "Systemmeldungen und Rueckfragen sind zentral steuerbar."),
    ]
    make_table(doc, ["Baustein", "Aktueller Stand", "Einordnung fuer Kollegen"], rows, [1.45, 3.05, 2.00])

    heading(doc, "Technische Dateien und Unterlagen", 2)
    for item in [
        "Frontend: app.js, styles.css und modulare Hilfsdateien fuer Router, State, Access Control, Review Workflow, SoA, Risk Engine, Versioning und Task Engine.",
        "Backend: backend_app.py plus database.py, storage.py, oidc_auth.py und malware_scan.py.",
        "Betrieb: OPERATIONS.md, Deployment-Quickstart, Go-Live-Checkliste, Backup-/Restore-Runbook, Docker Compose, nginx- und systemd-Vorlagen.",
        "Qualitaet: node --check, CSS-Klammerpruefung, Healthcheck und lokale Serverkontrolle per scripts/dev_server.sh.",
    ]:
        bullet(doc, item)


def add_demo_route(doc):
    heading(doc, "6. Empfohlene Demo-Reihenfolge", 1)
    rows = [
        ("1", "Dashboard", "Start mit Managementblick: Was muss als Naechstes passieren?"),
        ("2", "Kundenmodus", "Zeigen, wie ein Nicht-IT-Kunde durch Aufgaben und Nachweise gefuehrt wird."),
        ("3", "Projektstrukturplan", "Chronologische Arbeitspakete, Owner, Status und Verlinkung zu Dokumenten zeigen."),
        ("4", "Asset / Risiko / Vendor", "Gemeinsame Register als Datenbasis zeigen."),
        ("5", "Policy Management", "Policy-Paket und ausfuellbare Vorlagen oeffnen."),
        ("6", "Dokumentation", "Nachweis hochladen, verknuepfen, einreichen und Versionierung erklaeren."),
        ("7", "Beraterpruefung", "Pruefen, freigeben, Nacharbeit anfordern und Audit-Blocker zeigen."),
        ("8", "Auditpaket / Audit Trail", "Abschlusslogik und Nachvollziehbarkeit erklaeren."),
        ("9", "Team & Betrieb", "Einladungen, Rollen, Backups, Operations und Produktivbetrieb kurz einordnen."),
    ]
    make_table(doc, ["Nr.", "Station", "Vortragsfokus"], rows, [0.45, 1.55, 4.50])


def add_design_status(doc):
    heading(doc, "7. Design- und UX-Stand", 1)
    para(doc, "In den letzten Iterationen wurde die Plattform visuell deutlich in Richtung ruhiges B2B-SaaS gebracht.")
    for item in [
        "Weniger KI-/Demo-Look durch einheitlichere Buttons, Badges, Panels, Tabellen und Eingabefelder.",
        "Sidebar wurde beruhigt: klarere Gruppen, kleinere Badges, lesbarere aktive Zustaende.",
        "Widescreen-Layouts wurden stabilisiert, damit grosse Bildschirme nicht leer oder gequetscht wirken.",
        "Kundenmodus und Register wurden kundennaher aufgebaut: weniger Fachsprache, klarere naechste Schritte.",
        "Hilfetexte und Beispiele bleiben erhalten, damit Kunden wissen, was in Felder eingetragen werden soll.",
        "Die UI vermeidet Aussagen wie ISO-konform, NIS-2-erfuellt oder auditbereit ohne Beraterfreigabe.",
    ]:
        bullet(doc, item)


def add_limits_next(doc):
    heading(doc, "8. Grenzen und naechste Schritte", 1)
    callout(
        doc,
        "Realistischer Status",
        "Die Plattform ist vorfuehrbar und funktional breit. Fuer echten produktiven Kundenbetrieb braucht sie aber weiterhin kontrollierte QA, Security-Haertung, Datenschutzpruefung, Deployment-Konzept, Betriebshandbuch-Abnahme und ggf. Mandanten-/SSO-Integration.",
        fill=LIGHT_AMBER,
        border_color="E8C06F",
    )
    rows = [
        ("Design QA", "Alle Hauptseiten mit echten Beispielkunden, Desktop/Widescreen/Mobile pruefen.", "hoch"),
        ("Security Hardening", "Secrets, Cookies, HSTS, Upload-Scanner, Rollenrechte, MFA und Session-Konfiguration produktionsnah testen.", "hoch"),
        ("Backend-Betrieb", "PostgreSQL/S3/Reverse Proxy/Docker/systemd finalisieren und Restore-Tests dokumentieren.", "hoch"),
        ("Datenmodell", "Review-Status, Versionen, Audit Trail und Mandantenexporte gegen reale Workflows testen.", "mittel"),
        ("Kunden-Onboarding", "Einladungen, Rollen, Startprofil, Nacharbeit und Abgabecheck mit Pilotkunde durchspielen.", "hoch"),
        ("Recht/Datenschutz", "Datenschutz, AVV, Rollenverantwortung, Loeschkonzept und Aufbewahrung intern klaeren.", "hoch"),
        ("Content QA", "Policy-/Template-Texte fachlich pruefen und final freigeben lassen.", "mittel"),
        ("Produktstory", "Demo-Script, Screenshots und Grenzen einheitlich kommunizieren.", "mittel"),
    ]
    make_table(doc, ["Thema", "Naechste Arbeit", "Prioritaet"], rows, [1.45, 4.25, 0.80])


def add_appendix(doc):
    heading(doc, "9. Kurzer Abschluss fuer die Vorstellung", 1)
    para(
        doc,
        "Wenn ich es in einem Satz zusammenfasse: SFM Compliance soll Kunden nicht mit Normtext allein lassen, sondern sie durch die Vorbereitung fuehren und dem Berater eine saubere Pruef- und Freigabestrecke geben.",
        size=11.2,
        color=DARK,
        bold=True,
    )
    heading(doc, "10. Checkliste fuer die Live-Demo", 1)
    for item in [
        "Server starten und Healthcheck oeffnen: http://127.0.0.1:5173/api/health",
        "Browser hart neu laden, damit aktuelle Styles aktiv sind.",
        "Dashboard starten, danach Kundenmodus und Projektstrukturplan zeigen.",
        "Ein Nachweis-/Review-Beispiel zeigen, aber klar sagen: fachliche Freigabe bleibt beim Berater.",
        "Zum Schluss Backend-/Betrieb kurz zeigen: Rollen, Audit Trail, Backups, Healthcheck.",
    ]:
        bullet(doc, item)


def build():
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_talk_track(doc)
    add_function_map(doc)
    add_registers(doc)
    add_consultant_flow(doc)
    add_technical(doc)
    add_demo_route(doc)
    add_design_status(doc)
    add_limits_next(doc)
    add_appendix(doc)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
